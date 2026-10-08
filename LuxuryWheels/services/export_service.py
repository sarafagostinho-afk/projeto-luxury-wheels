# Módulo: export_service.py
# Descrição: Serviço centralizado para exportação de dados para CSV (com suporte a UTF-8/BOM) e Excel (.xlsx).


import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def exportar_para_csv(dados: list, caminho_ficheiro: str):
    # Exporta uma lista de dicionários para um ficheiro CSV, com codificação utf-8-sig para garantir a abertura correta de acentos no Microsoft Excel.
    if not dados:
        return False, "Não existem dados disponíveis para exportar."

    try:
        # Extrair as chaves do primeiro dicionário para servir de cabeçalho
        chaves = list(dados[0].keys())

        with open(caminho_ficheiro, mode="w", newline="", encoding="utf-8-sig") as f:
            escritor = csv.DictWriter(f, fieldnames=chaves)
            escritor.writeheader()
            for linha in dados:
                escritor.writerow(linha)

        return True, f"Ficheiro CSV exportado com sucesso para:\n{caminho_ficheiro}"
    except Exception as e:
        return False, f"Erro ao exportar ficheiro CSV: {str(e)}"


def exportar_para_excel(dados: list, caminho_ficheiro: str, titulo_folha: str = "Relatório"):
    # Exporta uma lista de dicionários para Excel (.xlsx).
    if not dados:
        return False, "Não existem dados disponíveis para exportar."

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = titulo_folha

        # Garantir visualização de grelha ativa
        ws.views.sheetView[0].showGridLines = True

        chaves = list(dados[0].keys())

        # Estilos visuais profissionais para o Excel
        fonte_cabecalho = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        fill_cabecalho = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        alinhamento_cabecalho = Alignment(horizontal="center", vertical="center")
        borda_fina = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB")
        )

        # Escrever cabeçalhos formatados
        ws.append([str(c).replace("_", " ").title() for c in chaves])
        for col_num in range(1, len(chaves) + 1):
            celula = ws.cell(row=1, column=col_num)
            celula.font = fonte_cabecalho
            celula.fill = fill_cabecalho
            celula.alignment = alinhamento_cabecalho
            celula.border = borda_fina

        ws.row_dimensions[1].height = 25

        # Escrever linhas de dados
        fonte_dados = Font(name="Arial", size=10)
        for linha_idx, item in enumerate(dados, start=2):
            valores = [item.get(c, "") for c in chaves]
            ws.append(valores)

            ws.row_dimensions[linha_idx].height = 20
            for col_num in range(1, len(chaves) + 1):
                celula = ws.cell(row=linha_idx, column=col_num)
                celula.font = fonte_dados
                celula.border = borda_fina
                celula.alignment = Alignment(
                    horizontal="left" if isinstance(celula.value, str) and len(str(celula.value)) > 15 else "center",
                    vertical="center")

        # Ajuste automático inteligente da largura das colunas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

        wb.save(caminho_ficheiro)
        return True, f"Ficheiro Excel exportado com sucesso para:\n{caminho_ficheiro}"
    except Exception as e:
        return False, f"Erro ao exportar ficheiro Excel: {str(e)}"